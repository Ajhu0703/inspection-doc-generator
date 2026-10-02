import os
import io
import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

st.set_page_config(page_title="材料驗收文件產生器", page_icon="📋", layout="wide")

st.title("📋 課程材料驗收文件自動產生器")
st.markdown("將材料 Excel 明細與照片自動整合填入 Word 範本中，產出標準驗收紀錄表。")

# 1. 基礎資訊輸入區
st.header("1. 填寫基本資訊")
col1, col2 = st.columns(2)
with col1:
    class_name = st.text_input("課程班級名稱與編號", value="115W0149-飲調輕食暨餐飲創業(桃園)-第1期")
with col2:
    delivery_date = st.text_input("進貨日期 (YYY/MM/DD)", value="115/02/03")

# 2. 檔案上傳區
st.header("2. 上傳驗收資料")
col_a, col_b = st.columns(2)

with col_a:
    excel_file = st.file_uploader("上傳材料明細 Excel (.xlsx)", type=["xlsx"])
    template_file = st.file_uploader("上傳空白 Word 範本 (.docx)", type=["docx"])

with col_b:
    image_files = st.file_uploader("上傳品項照片 (可複選多張照片)", type=["jpg", "png", "jpeg"], accept_multiple_files=True)

# 3. 核心處理邏輯
if st.button("🚀 開始生成驗收文件", type="primary"):
    if not excel_file:
        st.error("請上傳材料明細 Excel 檔案！")
    elif not template_file:
        st.error("請上傳 Word 空白範本！")
    else:
        try:
            # 讀取 Excel
            df = pd.read_excel(excel_file)
            
            # 建立照片檔名對映字典 (以檔名主體對應品名或項次)
            img_dict = {}
            if image_files:
                for img in image_files:
                    base_name = os.path.splitext(img.name)[0]
                    img_dict[base_name] = img.getvalue()

            # 載入 Word 範本
            doc = Document(template_file)

            # 填充班級與日期 (尋找並替換標籤，若有)
            for p in doc.paragraphs:
                if "課程班級" in p.text:
                    p.text = f"課程班級：{class_name}"
                if "進貨日期" in p.text:
                    p.text = f"進貨日期：{delivery_date}"

            # 假設表格為文件中的第一個表格
            if len(doc.tables) > 0:
                table = doc.tables[0]
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # 遍歷 Excel 的每一行資料並寫入表格
                for idx, row in df.iterrows():
                    # 新增一列
                    row_cells = table.add_row().cells
                    
                    # 取得資料 (預防欄位名稱差異)
                    item_no = str(row.get('品號', row.get('項次', idx + 1)))
                    item_name = str(row.get('品名', ''))
                    item_spec = str(row.get('規格', ''))
                    item_qty = f"{row.get('數量', '')} {row.get('單位', '')}".strip()
                    
                    # 填寫文字欄位
                    row_cells[0].text = item_no
                    row_cells[1].text = item_name
                    row_cells[2].text = item_spec
                    row_cells[3].text = item_qty
                    row_cells[5].text = "" # 備註欄保持空白
                    
                    # 置中文字
                    for c in [0, 3]:
                        for p in row_cells[c].paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    # 插入照片 (欄位 4)
                    cell_photo = row_cells[4]
                    p_photo = cell_photo.paragraphs[0]
                    p_photo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    # 比對照片檔名
                    matching_img = None
                    for key, img_bytes in img_dict.items():
                        if key in item_name or str(item_no) in key:
                            matching_img = img_bytes
                            break
                    
                    if matching_img:
                        image_stream = io.BytesIO(matching_img)
                        p_photo.add_run().add_picture(image_stream, width=Cm(3.5))
                    else:
                        p_photo.text = "（待補照片）"

            # 將結果寫入記憶體中的文件
            doc_io = io.BytesIO()
            doc.save(doc_io)
            doc_io.seek(0)

            st.success("🎉 驗收文件產製完成！")
            
            # 提供下載按鈕
            filename = f"驗收資料_{class_name}_{delivery_date}.docx".replace("/", "")
            st.download_button(
                label="📥 點此下載驗收文件 (.docx)",
                data=doc_io,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

        except Exception as e:
            st.error(f"處理檔案時發生錯誤：{str(e)}")